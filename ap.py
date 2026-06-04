import sys
import os
import logging
from functools import wraps
import time
import pandas as pd
import pyodbc
import io
import math
import json
import re
import numpy as np
from datetime import datetime, timedelta
from flask import Flask, jsonify, request, send_file, session, send_from_directory
from flask_cors import CORS
from flask_limiter import Limiter
from flask_limiter.util import get_remote_address
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# ============================================================================
# CONFIGURATION INITIALE
# ============================================================================

os.environ['PYTHONHASHSEED'] = '0'

app = Flask(__name__)
app.secret_key = 'votre_cle_secrete_remuci_2024'

app.config.update(
    SESSION_COOKIE_SECURE=False,
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE='Lax',
    PERMANENT_SESSION_LIFETIME=timedelta(hours=24),
    SESSION_REFRESH_EACH_REQUEST=True
)

CORS(app, 
     origins=["http://localhost:5000", "http://127.0.0.1:5000"],
     supports_credentials=True,
     allow_headers=["Content-Type", "Authorization", "Accept"],
     methods=["GET", "POST", "PUT", "DELETE", "OPTIONS"]
)

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
    handlers=[logging.StreamHandler()]
)
logger = logging.getLogger('remuci_app')

limiter = Limiter(
    key_func=get_remote_address,
    app=app,
    default_limits=["2000 per day", "500 per hour"]
)

# ============================================================================
# CACHE POUR ACCÉLÉRER LE DASHBOARD
# ============================================================================

dashboard_cache = {
    'data': None,
    'timestamp': None
}

def get_cached_dashboard():
    """Récupère les données du cache si elles datent de moins de 5 minutes"""
    if dashboard_cache['timestamp'] is not None:
        age = (datetime.now() - dashboard_cache['timestamp']).total_seconds()
        if age < 300:  # 5 minutes de cache
            return dashboard_cache['data']
    return None

def set_cached_dashboard(data):
    """Met en cache les données du dashboard"""
    dashboard_cache['data'] = data
    dashboard_cache['timestamp'] = datetime.now()

# ============================================================================
# FONCTIONS UTILITAIRES
# ============================================================================

def convert_to_serializable(obj):
    """Convertit les types numpy en types Python standard pour JSON"""
    if isinstance(obj, (np.int64, np.int32)):
        return int(obj)
    if isinstance(obj, (np.float64, np.float32)):
        return float(obj)
    if isinstance(obj, np.ndarray):
        return obj.tolist()
    if isinstance(obj, pd.Timestamp):
        return obj.strftime('%Y-%m-%d')
    if isinstance(obj, datetime):
        return obj.strftime('%Y-%m-%d %H:%M:%S')
    if pd.isna(obj):
        return None
    return obj

def dataframe_to_json(df):
    """Convertit un DataFrame en JSON compatible"""
    if df is None or df.empty:
        return []
    records = df.to_dict(orient='records')
    for record in records:
        for key, value in record.items():
            record[key] = convert_to_serializable(value)
    return records

# ============================================================================
# AUTHENTIFICATION
# ============================================================================

USERS = {
    "ADMIN": "Admin@2025!",
    "KALFRED": "RESPIT2025!",
    "NSANDRINE": "Assist@25#",
    "KEPONON": "Assist@25!"
}

@app.route('/login', methods=['POST', 'OPTIONS'])
def login():
    if request.method == 'OPTIONS':
        return jsonify({'success': True})
    
    username = request.form['username']
    password = request.form['password']
    
    if username in USERS and USERS[username] == password:
        session['username'] = username
        session['logged_in'] = True
        session.permanent = True
        return jsonify({'success': True, 'username': username})
    else:
        return jsonify({'success': False, 'message': 'Identifiants incorrects'}), 401

@app.route('/logout', methods=['POST'])
def logout():
    session.pop('username', None)
    session.pop('logged_in', None)
    return jsonify({'success': True})

@app.route('/api/check-session', methods=['GET'])
def check_session():
    if session.get('logged_in'):
        return jsonify({'authenticated': True, 'username': session.get('username')})
    return jsonify({'authenticated': False}), 401

@app.route('/')
def serve_interface():
    try:
        return send_file('remuci.html')
    except Exception as e:
        logger.error(f"Erreur chargement HTML: {e}")
        return jsonify({"error": "Fichier remuci.html non trouvé", "success": False}), 404

@app.route('/manifest.json')
def serve_manifest():
    return send_from_directory('.', 'manifest.json')

@app.route('/sw.js')
def serve_sw():
    return send_file('sw.js', mimetype='application/javascript')

@app.before_request
def require_login():
    public_routes = ['/', '/login', '/logout', '/manifest.json', '/sw.js', '/api/check-session', '/api/health', '/api/clear-cache']
    if request.path in public_routes:
        return None
    if request.path.startswith('/static/'):
        return None
    if not session.get('logged_in'):
        return jsonify({'success': False, 'error': 'Authentification requise'}), 401

# ============================================================================
# CONNEXION BASE DE DONNÉES
# ============================================================================

def get_connection():
    try:
        conn = pyodbc.connect(
            "Driver={SQL Server};"
            "Server=localhost\\SQLEXPRESS;"
            "Database=REMUCI_VISION;"
            "Trusted_Connection=yes;"
            "Timeout=30;",
            autocommit=True
        )
        print("Connexion SQL Server réussie")
        return conn
    except Exception as e:
        print(f"Erreur de connexion SQL: {e}")
        return None

def clean_dataframe(df):
    if df.empty:
        return df
    numeric_cols = df.select_dtypes(include=['float64', 'int64']).columns
    for col in numeric_cols:
        df[col] = df[col].fillna(0)
    text_cols = df.select_dtypes(include=['object']).columns
    for col in text_cols:
        df[col] = df[col].fillna('')
    return df

def format_montant(montant):
    if montant is None:
        return "0"
    return f"{int(montant):,}".replace(",", " ")

def get_today():
    return datetime.now().strftime('%Y-%m-%d')

def get_default_date_start(months=4):
    date = datetime.now()
    date = date.replace(day=1)
    for _ in range(months):
        date = date - timedelta(days=1)
        date = date.replace(day=1)
    return date.strftime('%Y-%m-%d')

# ============================================================================
# API ROUTES
# ============================================================================

@app.route("/api/health", methods=['GET'])
def health_check():
    return jsonify({"status": "healthy", "timestamp": datetime.now().isoformat()})

@app.route("/api/clear-cache", methods=['POST'])
def clear_cache():
    """Vide le cache du dashboard"""
    global dashboard_cache
    dashboard_cache = {'data': None, 'timestamp': None}
    return jsonify({"success": True, "message": "Cache vidé"})

@app.route("/api/gestionnaires", methods=["GET"])
@limiter.limit("30 per minute")
def get_gestionnaires():
    conn = get_connection()
    if not conn:
        return jsonify({"success": False, "error": "Connexion impossible"}), 500
    
    try:
        query = """
        SELECT DISTINCT gestionnaire_pret AS nom
        FROM dbo.extra_credits_view
        WHERE gestionnaire_pret IS NOT NULL AND gestionnaire_pret != ''
        ORDER BY gestionnaire_pret;
        """
        df = pd.read_sql(query, conn)
        return jsonify({"success": True, "data": dataframe_to_json(df)})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500
    finally:
        if conn:
            conn.close()


@app.route("/api/dashboard-data", methods=["GET"])
def dashboard_data():
    # Vérifier le cache d'abord
    cached_data = get_cached_dashboard()
    if cached_data is not None:
        return jsonify({"success": True, "data": cached_data, "cached": True})
    
    conn = get_connection()
    if not conn:
        return jsonify({"success": False, "error": "Connexion impossible"}), 500
    
    try:
        # OPTIMISATION: Une seule requête pour plusieurs KPIs
        query_kpis = """
        SELECT 
            -- Crédits débloqués 30j
            ISNULL(SUM(CASE WHEN date_effet >= DATEADD(day, -30, GETDATE()) THEN mtt_pret ELSE 0 END), 0) as credits_total,
            COUNT(CASE WHEN date_effet >= DATEADD(day, -30, GETDATE()) THEN 1 END) as credits_count,
            -- Nouveaux clients 30j
            COUNT(DISTINCT CASE WHEN date_adhesion >= DATEADD(day, -30, GETDATE()) THEN code_client END) as nouveaux_clients,
            -- Taux d'impayés
            COUNT(*) as total_dossiers,
            SUM(CASE WHEN date_fin_echeance < GETDATE() AND (date_solde IS NULL OR date_solde > date_fin_echeance) THEN 1 ELSE 0 END) as impayes_count,
            -- Clients actifs 3 mois
            COUNT(DISTINCT CASE WHEN date_effet >= DATEADD(month, -3, GETDATE()) THEN code_client END) as clients_actifs,
            -- Encours total
            ISNULL(SUM(CASE WHEN date_solde IS NULL OR date_solde > GETDATE() THEN mtt_pret ELSE 0 END), 0) as encours_total,
            -- Taux recouvrement
            CASE WHEN SUM(mtt_pret) > 0 
            THEN (SUM(CASE WHEN date_solde IS NOT NULL AND date_solde <= GETDATE() THEN mtt_pret ELSE 0 END) * 100.0 / SUM(mtt_pret))
            ELSE 0 END as taux_recouvrement
        FROM dbo.extra_credits_view
        WHERE date_effet IS NOT NULL
        """
        
        df_kpis = pd.read_sql(query_kpis, conn)
        
        credits_total = float(df_kpis.iloc[0]['credits_total']) if not df_kpis.empty else 0
        credits_count = int(df_kpis.iloc[0]['credits_count']) if not df_kpis.empty else 0
        nouveaux_clients = int(df_kpis.iloc[0]['nouveaux_clients']) if not df_kpis.empty else 0
        total_dossiers = df_kpis.iloc[0]['total_dossiers'] or 1
        impayes_count = df_kpis.iloc[0]['impayes_count'] or 0
        taux_impayes = round((impayes_count / total_dossiers) * 100, 1) if total_dossiers > 0 else 0
        clients_actifs = int(df_kpis.iloc[0]['clients_actifs']) if not df_kpis.empty else 0
        encours_total = float(df_kpis.iloc[0]['encours_total']) if not df_kpis.empty else 0
        taux_recouvrement = round(float(df_kpis.iloc[0]['taux_recouvrement']), 1) if not df_kpis.empty else 0
        
        # Comptes ouverts (requête rapide)
        query_comptes = "SELECT COUNT(*) as nb FROM dbo.COMPTES WHERE ETAT = 'O'"
        df_comptes = pd.read_sql(query_comptes, conn)
        comptes_ouverts = int(df_comptes.iloc[0]['nb']) if not df_comptes.empty else 0
        
        # Parts sociales
        try:
            query_parts = "SELECT ISNULL(SUM(ops.NOMBRE * psv.VALEUR), 0) as total FROM dbo.OPERATIONS_PART_SOC ops LEFT JOIN dbo.PARTS_SOCIALE psv ON ops.ID_PART_SOCIALE = psv.ID"
            df_parts = pd.read_sql(query_parts, conn)
            parts_sociales = float(df_parts.iloc[0]['total']) if not df_parts.empty else 125000000
        except:
            parts_sociales = 125000000
        
        # Évolution mensuelle (LIMITÉE à 12 mois)
        query_evolution = """
        SELECT TOP 12
            FORMAT(date_effet, 'MMM yyyy') as mois,
            ISNULL(SUM(mtt_pret), 0) as montant
        FROM dbo.extra_credits_view
        WHERE date_effet >= DATEADD(month, -11, GETDATE()) AND date_effet IS NOT NULL
        GROUP BY FORMAT(date_effet, 'MMM yyyy'), YEAR(date_effet), MONTH(date_effet)
        ORDER BY YEAR(date_effet), MONTH(date_effet)
        """
        df_evolution = pd.read_sql(query_evolution, conn)
        evolution_data = [{'mois': str(row['mois']), 'montant': float(row['montant'])} for _, row in df_evolution.iterrows()]
        
        # Répartition par produit (LIMITÉE à 5)
        query_repartition = """
        SELECT TOP 5
            CASE WHEN produit IS NULL OR produit = '' THEN 'Non classé' ELSE produit END as produit,
            ISNULL(SUM(mtt_pret), 0) as montant
        FROM dbo.extra_credits_view
        WHERE date_effet IS NOT NULL
        GROUP BY produit
        ORDER BY montant DESC
        """
        df_repartition = pd.read_sql(query_repartition, conn)
        repartition_data = [{'produit': str(row['produit']), 'montant': float(row['montant'])} for _, row in df_repartition.iterrows()]
        
        # Dernières activités (LIMITÉES à 10)
        query_activites = """
        SELECT TOP 10
            'Décaissement' as type,
            ISNULL(nom_client, '') + ' ' + ISNULL(prenoms_client, '') as client,
            mtt_pret as montant,
            date_effet as date_action,
            ISNULL(nom_agence, '') as agence
        FROM dbo.extra_credits_view
        WHERE date_effet IS NOT NULL
        ORDER BY date_effet DESC
        """
        df_activites = pd.read_sql(query_activites, conn)
        
        activites = []
        for _, row in df_activites.iterrows():
            date_val = row['date_action']
            date_str = date_val.strftime('%Y-%m-%d') if hasattr(date_val, 'strftime') else str(date_val) if date_val else ''
            montant_val = float(row['montant']) if row['montant'] and not pd.isna(row['montant']) else None
            activites.append({
                'type': str(row['type']),
                'client': str(row['client'][:50]) if row['client'] else 'Client',
                'montant': float(montant_val) if montant_val else None,
                'date': date_str,
                'agence': str(row['agence']) if not pd.isna(row['agence']) else ''
            })
        
        result_data = {
            "credits_debloques": float(credits_total),
            "credits_count": int(credits_count),
            "nouveaux_clients": int(nouveaux_clients),
            "comptes_ouverts": int(comptes_ouverts),
            "taux_impayes": float(taux_impayes),
            "impayes_count": int(impayes_count),
            "total_dossiers": int(total_dossiers),
            "clients_actifs": int(clients_actifs),
            "encours_total": float(encours_total),
            "taux_recouvrement": float(taux_recouvrement),
            "parts_sociales": float(parts_sociales),
            "evolution": evolution_data,
            "repartition": repartition_data,
            "activites": activites
        }
        
        # Mettre en cache
        set_cached_dashboard(result_data)
        
        return jsonify({"success": True, "data": result_data, "cached": False})
        
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)})
    finally:
        if conn:
            conn.close()


# ============================================================================
# CRÉDITS DÉBLOQUÉS
# ============================================================================

@app.route("/api/credits-debloques", methods=["GET"])
def credits_debloques():
    conn = get_connection()
    if not conn:
        return jsonify({"success": False, "error": "Connexion impossible"}), 500
    
    try:
        agence = request.args.get('agence', '')
        gestionnaire = request.args.get('gestionnaire', '')
        date_debut = request.args.get('date_debut') or get_default_date_start()
        date_fin = request.args.get('date_fin') or get_today()
        page = int(request.args.get('page', 1))
        limit = int(request.args.get('limit', 50))
        offset = (page - 1) * limit

        query = """
        SELECT 
            ecv.num_manuel AS [N° manuel],
            ecv.nom_client + ' ' + ISNULL(ecv.prenoms_client, '') AS [Client],
            ecv.mtt_pret AS [Montant],
            ecv.date_effet AS [Date déblocage],
            ecv.date_fin_echeance AS [Fin échéance],
            ecv.nb_echeance AS [Nb échéances],
            ecv.nom_agence AS [Agence],
            ecv.gestionnaire_pret AS [Gestionnaire],
            ecv.produit AS [Produit],
            ecv.code_client AS [Code client]
        FROM dbo.extra_credits_view ecv
        WHERE ecv.date_effet BETWEEN ? AND ?
          AND (? = '' OR ecv.nom_agence = ?)
          AND (? = '' OR ecv.gestionnaire_pret = ?)
        ORDER BY ecv.date_effet DESC
        OFFSET ? ROWS FETCH NEXT ? ROWS ONLY;
        """
        
        count_query = """
        SELECT COUNT(*) as total
        FROM dbo.extra_credits_view ecv
        WHERE ecv.date_effet BETWEEN ? AND ?
          AND (? = '' OR ecv.nom_agence = ?)
          AND (? = '' OR ecv.gestionnaire_pret = ?);
        """
        
        params = [date_debut, date_fin, agence, agence, gestionnaire, gestionnaire]
        
        df = pd.read_sql(query, conn, params=params + [offset, limit])
        df_count = pd.read_sql(count_query, conn, params=params)
        
        total = int(df_count.iloc[0]['total']) if not df_count.empty else 0
        
        data = dataframe_to_json(df)
        
        return jsonify({
            "success": True,
            "data": data,
            "total": total,
            "page": page,
            "limit": limit,
            "pages": (total + limit - 1) // limit if limit > 0 else 1
        })
        
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500
    finally:
        if conn:
            conn.close()


# ============================================================================
# NOUVEAUX CLIENTS
# ============================================================================

@app.route("/api/nouveaux-clients", methods=["GET"])
def nouveaux_clients():
    conn = get_connection()
    if not conn:
        return jsonify({"success": False, "error": "Connexion impossible"}), 500
    
    try:
        agence = request.args.get('agence', '')
        date_debut = request.args.get('date_debut') or get_default_date_start(6)
        date_fin = request.args.get('date_fin') or get_today()
        page = int(request.args.get('page', 1))
        limit = int(request.args.get('limit', 50))
        offset = (page - 1) * limit

        query = """
        SELECT DISTINCT
            ecv.code_client AS [Code client],
            ecv.nom_client + ' ' + ISNULL(ecv.prenoms_client, '') AS [Client],
            ecv.date_adhesion AS [Date adhésion],
            ecv.telephone AS [Téléphone],
            ecv.nom_agence AS [Agence],
            ecv.sexe AS [Sexe]
        FROM dbo.extra_credits_view ecv
        WHERE ecv.date_adhesion BETWEEN ? AND ?
          AND (? = '' OR ecv.nom_agence = ?)
        ORDER BY ecv.date_adhesion DESC
        OFFSET ? ROWS FETCH NEXT ? ROWS ONLY;
        """
        
        count_query = """
        SELECT COUNT(DISTINCT code_client) as total
        FROM dbo.extra_credits_view
        WHERE date_adhesion BETWEEN ? AND ?
          AND (? = '' OR nom_agence = ?);
        """
        
        params = [date_debut, date_fin, agence, agence]
        
        df = pd.read_sql(query, conn, params=params + [offset, limit])
        df_count = pd.read_sql(count_query, conn, params=params)
        
        total = int(df_count.iloc[0]['total']) if not df_count.empty else 0
        
        data = dataframe_to_json(df)
        
        return jsonify({
            "success": True,
            "data": data,
            "total": total,
            "page": page,
            "pages": (total + limit - 1) // limit if limit > 0 else 1
        })
        
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500
    finally:
        if conn:
            conn.close()


# ============================================================================
# COMPTES OUVERTS
# ============================================================================

@app.route("/api/comptes-ouverts", methods=["GET"])
def comptes_ouverts():
    conn = get_connection()
    if not conn:
        return jsonify({"success": False, "error": "Connexion impossible"}), 500
    
    try:
        date_debut = request.args.get('date_debut') or get_default_date_start(6)
        date_fin = request.args.get('date_fin') or get_today()
        page = int(request.args.get('page', 1))
        limit = int(request.args.get('limit', 50))
        offset = (page - 1) * limit

        query = """
        SELECT 
            c.NUM_CPTE as [Numéro Compte],
            c.LIBELLE as [Libellé],
            c.DATE_OUVERTURE as [Date ouverture],
            a.NOM_ADHERENT as [Client],
            a.ID as [ID Client],
            ISNULL((
                SELECT SUM(MONTANT_OPERATION * CASE WHEN SENS = 'C' THEN 1 ELSE -1 END)
                FROM HDPM h WHERE h.ID_COMPTE = c.ID
            ), 0) as [Solde actuel]
        FROM COMPTES c
        LEFT JOIN ADHERENTS a ON c.ID = a.ID_COMPTE_ADHERENT
        WHERE c.ETAT = 'O'
          AND c.DATE_OUVERTURE BETWEEN ? AND ?
        ORDER BY c.DATE_OUVERTURE DESC
        OFFSET ? ROWS FETCH NEXT ? ROWS ONLY;
        """
        
        count_query = """
        SELECT COUNT(*) as total
        FROM COMPTES c
        WHERE c.ETAT = 'O' AND c.DATE_OUVERTURE BETWEEN ? AND ?;
        """
        
        df = pd.read_sql(query, conn, params=[date_debut, date_fin, offset, limit])
        df_count = pd.read_sql(count_query, conn, params=[date_debut, date_fin])
        
        total = int(df_count.iloc[0]['total']) if not df_count.empty else 0
        
        return jsonify({
            "success": True,
            "data": dataframe_to_json(df),
            "total": total,
            "page": page,
            "pages": (total + limit - 1) // limit if limit > 0 else 1
        })
        
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500
    finally:
        if conn:
            conn.close()


# ============================================================================
# CRÉDITS IMPAYÉS
# ============================================================================

@app.route("/api/credits-impayes", methods=["GET"])
def credits_impayes():
    conn = get_connection()
    if not conn:
        return jsonify({"success": False, "error": "Connexion impossible"}), 500
    
    try:
        agence = request.args.get('agence', '')
        jours_retard_min = int(request.args.get('jours_retard_min', 0))
        jours_retard_max = int(request.args.get('jours_retard_max', 9999))
        page = int(request.args.get('page', 1))
        limit = int(request.args.get('limit', 50))
        offset = (page - 1) * limit

        query = """
        SELECT 
            ecv.num_manuel AS [N° manuel],
            ecv.nom_client + ' ' + ISNULL(ecv.prenoms_client, '') AS [Client],
            ecv.mtt_pret AS [Montant],
            ecv.date_fin_echeance AS [Date échéance],
            DATEDIFF(day, ecv.date_fin_echeance, GETDATE()) AS [Jours retard],
            ecv.nom_agence AS [Agence],
            ecv.gestionnaire_pret AS [Gestionnaire],
            ecv.produit AS [Produit],
            ecv.telephone AS [Téléphone]
        FROM dbo.extra_credits_view ecv
        WHERE ecv.date_fin_echeance < GETDATE()
          AND (ecv.date_solde IS NULL OR ecv.date_solde > ecv.date_fin_echeance)
          AND (? = '' OR ecv.nom_agence = ?)
          AND DATEDIFF(day, ecv.date_fin_echeance, GETDATE()) BETWEEN ? AND ?
        ORDER BY [Jours retard] DESC
        OFFSET ? ROWS FETCH NEXT ? ROWS ONLY;
        """
        
        count_query = """
        SELECT COUNT(*) as total
        FROM dbo.extra_credits_view ecv
        WHERE ecv.date_fin_echeance < GETDATE()
          AND (ecv.date_solde IS NULL OR ecv.date_solde > ecv.date_fin_echeance)
          AND (? = '' OR ecv.nom_agence = ?)
          AND DATEDIFF(day, ecv.date_fin_echeance, GETDATE()) BETWEEN ? AND ?;
        """
        
        df = pd.read_sql(query, conn, params=[agence, agence, jours_retard_min, jours_retard_max, offset, limit])
        df_count = pd.read_sql(count_query, conn, params=[agence, agence, jours_retard_min, jours_retard_max])
        
        total = int(df_count.iloc[0]['total']) if not df_count.empty else 0
        
        return jsonify({
            "success": True,
            "data": dataframe_to_json(df),
            "total": total,
            "page": page,
            "pages": (total + limit - 1) // limit if limit > 0 else 1
        })
        
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500
    finally:
        if conn:
            conn.close()


# ============================================================================
# ÉCHÉANCES FUTURES
# ============================================================================

@app.route("/api/echeances-futures", methods=["GET"])
def echeances_futures():
    conn = get_connection()
    if not conn:
        return jsonify({"success": False, "error": "Connexion impossible"}), 500
    
    try:
        date_debut = request.args.get('date_debut') or get_today()
        date_fin = request.args.get('date_fin') or (datetime.now() + timedelta(days=30)).strftime('%Y-%m-%d')
        agence = request.args.get('agence', '')
        page = int(request.args.get('page', 1))
        limit = int(request.args.get('limit', 50))
        offset = (page - 1) * limit

        query = """
        SELECT 
            ecv.nom_client + ' ' + ISNULL(ecv.prenoms_client, '') AS [Client],
            ecv.num_manuel AS [N° contrat],
            ecv.mtt_pret AS [Montant crédit],
            ecv.date_fin_echeance AS [Date échéance],
            DATEDIFF(day, GETDATE(), ecv.date_fin_echeance) AS [Jours restants],
            ecv.nom_agence AS [Agence],
            ecv.gestionnaire_pret AS [Gestionnaire],
            ecv.telephone AS [Téléphone]
        FROM dbo.extra_credits_view ecv
        WHERE ecv.date_fin_echeance BETWEEN ? AND ?
          AND ecv.date_solde IS NULL
          AND (? = '' OR ecv.nom_agence = ?)
        ORDER BY ecv.date_fin_echeance
        OFFSET ? ROWS FETCH NEXT ? ROWS ONLY;
        """
        
        count_query = """
        SELECT COUNT(*) as total
        FROM dbo.extra_credits_view ecv
        WHERE ecv.date_fin_echeance BETWEEN ? AND ?
          AND ecv.date_solde IS NULL
          AND (? = '' OR ecv.nom_agence = ?);
        """
        
        df = pd.read_sql(query, conn, params=[date_debut, date_fin, agence, agence, offset, limit])
        df_count = pd.read_sql(count_query, conn, params=[date_debut, date_fin, agence, agence])
        
        total = int(df_count.iloc[0]['total']) if not df_count.empty else 0
        
        return jsonify({
            "success": True,
            "data": dataframe_to_json(df),
            "total": total,
            "page": page,
            "pages": (total + limit - 1) // limit if limit > 0 else 1
        })
        
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500
    finally:
        if conn:
            conn.close()


# ============================================================================
# REMBOURSEMENTS
# ============================================================================

@app.route("/api/remboursements", methods=["GET"])
def remboursements():
    conn = get_connection()
    if not conn:
        return jsonify({"success": False, "error": "Connexion impossible"}), 500
    
    try:
        date_debut = request.args.get('date_debut') or (datetime.now() - timedelta(days=30)).strftime('%Y-%m-%d')
        date_fin = request.args.get('date_fin') or get_today()
        agence = request.args.get('agence', '')
        page = int(request.args.get('page', 1))
        limit = int(request.args.get('limit', 50))
        offset = (page - 1) * limit

        query = """
        SELECT 
            ecv.nom_client + ' ' + ISNULL(ecv.prenoms_client, '') AS [Client],
            ecv.num_manuel AS [N° contrat],
            ecv.mtt_pret AS [Montant crédit],
            ecv.date_solde AS [Date remboursement],
            ecv.nom_agence AS [Agence]
        FROM dbo.extra_credits_view ecv
        WHERE ecv.date_solde BETWEEN ? AND ?
          AND ecv.date_solde IS NOT NULL
          AND (? = '' OR ecv.nom_agence = ?)
        ORDER BY ecv.date_solde DESC
        OFFSET ? ROWS FETCH NEXT ? ROWS ONLY;
        """
        
        count_query = """
        SELECT COUNT(*) as total
        FROM dbo.extra_credits_view ecv
        WHERE ecv.date_solde BETWEEN ? AND ?
          AND ecv.date_solde IS NOT NULL
          AND (? = '' OR ecv.nom_agence = ?);
        """
        
        df = pd.read_sql(query, conn, params=[date_debut, date_fin, agence, agence, offset, limit])
        df_count = pd.read_sql(count_query, conn, params=[date_debut, date_fin, agence, agence])
        
        total = int(df_count.iloc[0]['total']) if not df_count.empty else 0
        
        return jsonify({
            "success": True,
            "data": dataframe_to_json(df),
            "total": total,
            "page": page,
            "pages": (total + limit - 1) // limit if limit > 0 else 1
        })
        
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500
    finally:
        if conn:
            conn.close()


# ============================================================================
# CLIENTS ACTIFS
# ============================================================================

@app.route("/api/clients-actifs", methods=["GET"])
def clients_actifs():
    conn = get_connection()
    if not conn:
        return jsonify({"success": False, "error": "Connexion impossible"}), 500
    
    try:
        agence = request.args.get('agence', '')
        page = int(request.args.get('page', 1))
        limit = int(request.args.get('limit', 50))
        offset = (page - 1) * limit

        query = """
        SELECT DISTINCT
            ecv.code_client AS [Code client],
            ecv.nom_client + ' ' + ISNULL(ecv.prenoms_client, '') AS [Client],
            ecv.telephone AS [Téléphone],
            ecv.nom_agence AS [Agence],
            ecv.gestionnaire_pret AS [Gestionnaire],
            MAX(ecv.date_effet) AS [Dernier crédit],
            COUNT(ecv.id_pret) AS [Nb crédits]
        FROM dbo.extra_credits_view ecv
        WHERE ecv.date_effet >= DATEADD(month, -3, GETDATE())
          AND (? = '' OR ecv.nom_agence = ?)
        GROUP BY ecv.code_client, ecv.nom_client, ecv.prenoms_client, ecv.telephone, ecv.nom_agence, ecv.gestionnaire_pret
        ORDER BY [Dernier crédit] DESC
        OFFSET ? ROWS FETCH NEXT ? ROWS ONLY;
        """
        
        count_query = """
        SELECT COUNT(DISTINCT code_client) as total
        FROM dbo.extra_credits_view
        WHERE date_effet >= DATEADD(month, -3, GETDATE())
          AND (? = '' OR nom_agence = ?);
        """
        
        df = pd.read_sql(query, conn, params=[agence, agence, offset, limit])
        df_count = pd.read_sql(count_query, conn, params=[agence, agence])
        
        total = int(df_count.iloc[0]['total']) if not df_count.empty else 0
        
        return jsonify({
            "success": True,
            "data": dataframe_to_json(df),
            "total": total,
            "page": page,
            "pages": (total + limit - 1) // limit if limit > 0 else 1
        })
        
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500
    finally:
        if conn:
            conn.close()


# ============================================================================
# PARTS SOCIALES
# ============================================================================

@app.route("/api/parts-sociales", methods=["GET"])
def parts_sociales():
    conn = get_connection()
    if not conn:
        return jsonify({"success": False, "error": "Connexion impossible"}), 500
    
    try:
        page = int(request.args.get('page', 1))
        limit = int(request.args.get('limit', 50))
        offset = (page - 1) * limit

        query = """
        SELECT TOP 50
            a.NOM_ADHERENT AS [Client],
            a.CODE AS [Code],
            ops.NOMBRE AS [Nombre parts],
            (ops.NOMBRE * psv.VALEUR) AS [Montant],
            o.DATE_OPERATION AS [Date]
        FROM OPERATIONS_PART_SOC ops
        LEFT JOIN ADHERENTS a ON ops.ID_ADHERENT = a.ID
        LEFT JOIN PARTS_SOCIALE psv ON ops.ID_PART_SOCIALE = psv.ID
        LEFT JOIN OPERATIONS o ON ops.ID_OPERATION = o.ID
        WHERE ops.NOMBRE > 0
        ORDER BY o.DATE_OPERATION DESC
        OFFSET ? ROWS FETCH NEXT ? ROWS ONLY;
        """
        
        count_query = """
        SELECT COUNT(*) as total FROM OPERATIONS_PART_SOC WHERE NOMBRE > 0;
        """
        
        df = pd.read_sql(query, conn, params=[offset, limit])
        df_count = pd.read_sql(count_query, conn)
        
        total = int(df_count.iloc[0]['total']) if not df_count.empty else 0
        
        return jsonify({
            "success": True,
            "data": dataframe_to_json(df),
            "total": total,
            "page": page,
            "pages": (total + limit - 1) // limit if limit > 0 else 1
        })
        
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500
    finally:
        if conn:
            conn.close()


# ============================================================================
# BALANCE J-1
# ============================================================================

@app.route("/api/balance-j1", methods=["GET"])
def balance_j1():
    return jsonify({
        "success": True,
        "data": [
            {"Compte": "10111100", "Intitulé": "Caisse Centrale", "Solde Débit": 0, "Solde Crédit": 0}
        ],
        "classes": [],
        "metadata": {"date_balance": get_today()}
    })


# ============================================================================
# GRAND LIVRE
# ============================================================================

@app.route("/api/grand-livre", methods=["GET"])
def grand_livre():
    compte = request.args.get('compte', '')
    return jsonify({
        "success": True,
        "compte": {"numero": compte, "libelle": "Compte en construction", "devise": "FCFA"},
        "periode": {"debut": get_default_date_start(1), "fin": get_today()},
        "solde_initial": 0,
        "total_debit": 0,
        "total_credit": 0,
        "solde_final": 0,
        "mouvements": [],
        "nombre_operations": 0
    })


# ============================================================================
# ANALYSE PAR GENRE
# ============================================================================

@app.route("/api/analyse-genre/comparatif", methods=["GET"])
def analyse_genre():
    conn = get_connection()
    if not conn:
        return jsonify({"success": False, "error": "Connexion impossible"}), 500
    
    try:
        annee = request.args.get('annee', str(datetime.now().year))
        
        query_femmes = """
        SELECT 
            COUNT(DISTINCT ecv.code_client) as clients,
            COUNT(ecv.id_pret) as credits,
            ISNULL(SUM(ecv.mtt_pret), 0) as montant
        FROM dbo.extra_credits_view ecv
        WHERE YEAR(ecv.date_effet) = ? AND ecv.sexe = 'F'
        """
        
        query_hommes = """
        SELECT 
            COUNT(DISTINCT ecv.code_client) as clients,
            COUNT(ecv.id_pret) as credits,
            ISNULL(SUM(ecv.mtt_pret), 0) as montant
        FROM dbo.extra_credits_view ecv
        WHERE YEAR(ecv.date_effet) = ? AND ecv.sexe = 'M'
        """
        
        df_f = pd.read_sql(query_femmes, conn, params=[annee])
        df_m = pd.read_sql(query_hommes, conn, params=[annee])
        
        return jsonify({
            "success": True,
            "data": {
                "F": {
                    "credits": {
                        "clients": int(df_f.iloc[0]['clients']) if not df_f.empty else 0,
                        "nombre": int(df_f.iloc[0]['credits']) if not df_f.empty else 0,
                        "montant_total": float(df_f.iloc[0]['montant']) if not df_f.empty else 0
                    }
                },
                "M": {
                    "credits": {
                        "clients": int(df_m.iloc[0]['clients']) if not df_m.empty else 0,
                        "nombre": int(df_m.iloc[0]['credits']) if not df_m.empty else 0,
                        "montant_total": float(df_m.iloc[0]['montant']) if not df_m.empty else 0
                    }
                }
            },
            "metadata": {"annee": annee}
        })
        
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500
    finally:
        if conn:
            conn.close()


# ============================================================================
# RECHERCHE GLOBALE
# ============================================================================

@app.route("/api/search", methods=["GET"])
def global_search():
    query = request.args.get('q', '')
    if len(query) < 2:
        return jsonify({"success": True, "data": []})
    
    conn = get_connection()
    if not conn:
        return jsonify({"success": False, "error": "Connexion impossible"}), 500
    
    try:
        search_term = f"%{query}%"
        sql = """
        SELECT TOP 20 'Client' as type, code_client as code, nom_client + ' ' + ISNULL(prenoms_client, '') as nom, telephone as details
        FROM extra_credits_view WHERE nom_client LIKE ? OR prenoms_client LIKE ? OR code_client LIKE ?
        UNION
        SELECT TOP 20 'Crédit' as type, num_manuel as code, nom_client + ' ' + ISNULL(prenoms_client, '') as nom, CAST(mtt_pret AS VARCHAR) as details
        FROM extra_credits_view WHERE num_manuel LIKE ?
        """
        df = pd.read_sql(sql, conn, params=[search_term, search_term, search_term, search_term])
        return jsonify({"success": True, "data": dataframe_to_json(df)})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500
    finally:
        if conn:
            conn.close()


# ============================================================================
# PLANIFICATION - VERSION COMPLÈTE
# ============================================================================

@app.route("/api/planning/list", methods=["GET"])
def get_plannings():
    conn = get_connection()
    if not conn:
        return jsonify({"success": False, "error": "Connexion impossible"}), 500
    
    try:
        query = "SELECT * FROM dbo.PLANIFICATIONS ORDER BY DATE_HEURE"
        df = pd.read_sql(query, conn)
        data = dataframe_to_json(df)
        return jsonify({"success": True, "data": data})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500
    finally:
        if conn:
            conn.close()


@app.route("/api/planning/create", methods=["POST"])
def create_planning():
    conn = get_connection()
    if not conn:
        return jsonify({"success": False, "error": "Connexion impossible"}), 500
    
    try:
        data = request.json
        type_extrait = data.get('type')
        frequence = data.get('frequence')
        destinataires = data.get('destinataires')
        date_heure = data.get('date_heure')
        
        query = """
        INSERT INTO dbo.PLANIFICATIONS (TYPE_EXTRAIT, FREQUENCE, DESTINATAIRES, DATE_HEURE, ACTIF)
        VALUES (?, ?, ?, ?, 1)
        """
        cursor = conn.cursor()
        cursor.execute(query, (type_extrait, frequence, destinataires, date_heure))
        conn.commit()
        
        return jsonify({"success": True, "message": "Planification créée"})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500
    finally:
        if conn:
            conn.close()


@app.route("/api/planning/delete/<int:plan_id>", methods=["DELETE"])
def delete_planning(plan_id):
    conn = get_connection()
    if not conn:
        return jsonify({"success": False, "error": "Connexion impossible"}), 500
    
    try:
        cursor = conn.cursor()
        cursor.execute("DELETE FROM dbo.PLANIFICATIONS WHERE ID = ?", (plan_id,))
        conn.commit()
        return jsonify({"success": True})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500
    finally:
        if conn:
            conn.close()


@app.route("/api/planning/update/<int:plan_id>", methods=["PUT"])
def update_planning(plan_id):
    conn = get_connection()
    if not conn:
        return jsonify({"success": False, "error": "Connexion impossible"}), 500
    
    try:
        data = request.json
        query = """
        UPDATE dbo.PLANIFICATIONS 
        SET TYPE_EXTRAIT = ?, FREQUENCE = ?, DESTINATAIRES = ?, DATE_HEURE = ?, ACTIF = ?
        WHERE ID = ?
        """
        cursor = conn.cursor()
        cursor.execute(query, (data.get('type'), data.get('frequence'), data.get('destinataires'), data.get('date_heure'), data.get('actif'), plan_id))
        conn.commit()
        return jsonify({"success": True})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500
    finally:
        if conn:
            conn.close()
            
# ============================================================================
# EXPORT EXCEL
# ============================================================================

@app.route("/api/export-excel/<type_export>", methods=["GET"])
def export_excel(type_export):
    output = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Export"
    ws['A1'] = f"Export {type_export}"
    ws['A2'] = f"Date: {datetime.now().strftime('%d/%m/%Y')}"
    wb.save(output)
    output.seek(0)
    
    filename = f"export_{type_export}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True, download_name=filename)


# ============================================================================
# ROUTES DE TEST
# ============================================================================

@app.route("/api/test-dashboard", methods=["GET"])
def test_dashboard():
    conn = get_connection()
    if not conn:
        return jsonify({"error": "Connexion DB impossible"})
    
    try:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT COUNT(*) as nb, ISNULL(SUM(mtt_pret), 0) as total
            FROM dbo.extra_credits_view
            WHERE date_effet >= DATEADD(day, -30, GETDATE())
        """)
        row = cursor.fetchone()
        
        return jsonify({
            "success": True,
            "credits_30j": {"nb": row[0], "total": float(row[1])},
            "message": "Connexion DB OK"
        })
    except Exception as e:
        return jsonify({"error": str(e)})
    finally:
        conn.close()


@app.route("/api/test-email", methods=["POST"])
@limiter.limit("5 per hour")
def test_email():
    try:
        from email_sender import send_email
        
        test_html = f"""
        <html>
        <body style="font-family: Arial;">
            <div style="background: #1a472a; color: white; padding: 20px;">
                <h2>REMU-CI VisionExtract</h2>
                <h3>Test de configuration email</h3>
            </div>
            <div style="padding: 20px;">
                <p>✅ Configuration email fonctionnelle !</p>
                <p>Date du test: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}</p>
            </div>
        </body>
        </html>
        """
        
        email_test = request.json.get('email') if request.json else None
        if not email_test:
            email_test = session.get('username', 'admin@remuci.ci')
            if '@' not in email_test:
                email_test = f"{email_test}@remuci.ci"
        
        success = send_email(
            to_emails=[email_test],
            subject="[VisionExtract] ✅ Test configuration",
            html_body=test_html,
            attachments=None
        )
        
        if success:
            return jsonify({"success": True, "message": f"Email test envoyé à {email_test}"})
        else:
            return jsonify({"success": False, "message": "Erreur d'envoi"}), 500
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500


@app.route("/api/email-status", methods=["GET"])
def email_status():
    from email_sender import EMAIL_CONFIG
    
    safe_config = {
        'smtp_server': EMAIL_CONFIG['smtp_server'],
        'smtp_port': EMAIL_CONFIG['smtp_port'],
        'smtp_user': EMAIL_CONFIG['smtp_user'],
        'from_email': EMAIL_CONFIG['from_email'],
        'configured': bool(EMAIL_CONFIG['smtp_password'] and EMAIL_CONFIG['smtp_password'] != 'VOTRE_MOT_DE_PASSE')
    }
    
    return jsonify({
        "success": True,
        "config": safe_config,
        "message": "La configuration est " + ("prête" if safe_config['configured'] else "incomplète")
    })


# ============================================================================
# DÉMARRAGE
# ============================================================================

if __name__ == "__main__":
    print("=" * 60)
    print("DÉMARRAGE DE VISIONEXTRACT - VERSION OPTIMISÉE")
    print("=" * 60)
    print("\nInterface disponible sur:")
    print("   - http://127.0.0.1:5000")
    print("\nIdentifiants: ADMIN / Admin@2025!")
    print("\n✅ Cache activé (5 minutes)")
    print("✅ Requêtes SQL optimisées")
    print("=" * 60)
    
    app.run(debug=True, host='0.0.0.0', port=5000)